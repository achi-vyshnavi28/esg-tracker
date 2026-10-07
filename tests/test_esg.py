"""Tests for the ESG reporting cycle: every planted data problem is caught, the GHG maths is right, and the outputs build."""
from collections import Counter

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from esg import calc, simulate, validate
from esg.run import run


@pytest.fixture(scope="module")
def cycle():
    sites, subs, bills, company, baseline = simulate.generate()
    clean, issues, recon = validate.validate(sites, subs, bills)
    return sites, subs, bills, company, baseline, clean, issues, recon


def test_every_planted_problem_is_found_exactly_once(cycle):
    *_, issues, recon = cycle
    assert Counter(i.check for i in issues) == {
        "missing_submission": 6, "duplicate_submission": 3, "unit_error": 2, "negative_value": 1,
        "recycled_exceeds_total": 2, "bill_variance": 3}


def test_clean_data_is_complete_and_valid(cycle):
    sites, _, _, _, _, clean, _, _ = cycle
    assert len(clean) == len(sites) * 12 == 492
    assert len({(r["site_id"], r["month"]) for r in clean}) == 492
    for r in clean:
        assert all(r[f] is not None and r[f] >= 0 for f in simulate.FIELDS)
        assert r["recycled_kg"] <= r["waste_kg"]


def test_electricity_matches_every_utility_bill_after_reconciliation(cycle):
    _, _, bills, _, _, clean, _, _ = cycle
    for b in bills:
        got = sum(r["grid_kwh"] for r in clean if r["site_id"] == b["site_id"] and simulate.quarter(r["month"]) == b["quarter"])
        assert got == pytest.approx(b["billed_kwh"], rel=1e-3)


def test_unit_error_is_converted_from_mwh(cycle):
    *_, clean, issues, _ = cycle
    for i in [x for x in issues if x.check == "unit_error"]:
        row = next(r for r in clean if r["site_id"] == i.site_id and r["month"] == i.period)
        assert row["grid_kwh"] > 1000


def test_emission_factor_maths():
    r = {"diesel_litres": 1000, "lpg_kg": 1000, "refrigerant_kg": 1, "grid_kwh": 1000}
    e = calc.row_emissions(r)
    assert e["scope1_diesel"] == pytest.approx(2.651) and e["scope1_lpg"] == pytest.approx(2.985)
    assert e["scope1_refrigerant"] == pytest.approx(1.924) and e["scope2_grid"] == pytest.approx(0.727)
    c = calc.company_emissions({"freight_tonne_km": 1_000_000, "air_passenger_km": 1000})
    assert c["scope3_freight"] == pytest.approx(75.47) and c["scope3_air"] == pytest.approx(0.27257)


def test_quarters_add_up_to_the_year(cycle):
    sites, _, _, company, baseline, clean, _, _ = cycle
    inv = calc.inventory(sites, clean, company, baseline)
    for k in ("scope1", "scope2", "scope3", "water_kl"):
        assert sum(inv["by_quarter"][q][k] for q in calc.QUARTERS) == pytest.approx(inv["year"][k])
    assert [k["met"] for k in inv["kpis"]] == [False, False, True, True]


def test_workbook_and_deck_are_built(tmp_path):
    s = run(tmp_path)
    wb = load_workbook(tmp_path / "esg_workbook.xlsx")
    assert wb.sheetnames == ["KPIs vs Targets", "Factors", "Activity", "Scope 3", "By Quarter", "Validation Log",
                             "Reconciliation"]
    assert wb["Activity"]["P2"].value.startswith("=(I2*Factors!$C$2")      # emissions are live formulas
    assert wb["Validation Log"].max_row - 1 == len(s["issues"]) == 17
    deck = Presentation(tmp_path / "quarterly_performance.pptx")
    assert len(deck.slides) == 7
    titles = [sl.shapes.title.text for sl in deck.slides]
    assert "Performance against targets" in titles and "Data quality: what was checked and fixed" in titles
