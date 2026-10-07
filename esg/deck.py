"""Quarterly sustainability performance deck for CXO review (PowerPoint, native editable charts)."""
from collections import Counter

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from esg.calc import QUARTERS

GREEN, RED, DARK = RGBColor(0x2F, 0x5D, 0x50), RGBColor(0xC0, 0x39, 0x2B), RGBColor(0x33, 0x33, 0x33)


def _slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    return s


def _bullets(slide, lines, top=1.5, size=16):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12), Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.font.size, p.font.color.rgb = line, Pt(size), DARK
        p.space_after = Pt(8)


def build(path, inv, issues, recon, n_submissions, n_expected):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    y, q = inv["year"], inv["by_quarter"]
    total = y["scope1"] + y["scope2"] + y["scope3"]

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly Sustainability Performance, FY 2025-26"
    s.placeholders[1].text = "ESG data review: 36 stores, 3 warehouses, 2 corporate offices (simulated network)"

    s = _slide(prs, "Performance against targets")
    rows = len(inv["kpis"]) + 1
    table = s.shapes.add_table(rows, 5, Inches(0.6), Inches(1.6), Inches(12), Inches(0.6) * rows).table
    for j, h in enumerate(["KPI", "FY 2024-25", "FY 2025-26", "Target", "Status"]):
        table.cell(0, j).text = h
    for i, k in enumerate(inv["kpis"], start=1):
        vals = [k["kpi"], "" if k["baseline"] is None else f"{k['baseline']:.2f}", f"{k['actual']:.2f}", k["target"],
                "Met" if k["met"] else "Not met"]
        for j, v in enumerate(vals):
            table.cell(i, j).text = v
        run = table.cell(i, 4).text_frame.paragraphs[0].runs[0]
        run.font.bold, run.font.color.rgb = True, GREEN if k["met"] else RED
    met = sum(k["met"] for k in inv["kpis"])
    _bullets(s, [f"{met} of {len(inv['kpis'])} targets met. Total footprint {total:,.0f} tCO2e "
                 f"(Scope 1 {y['scope1']:,.0f}, Scope 2 {y['scope2']:,.0f}, Scope 3 {y['scope3']:,.0f})."], top=5.2)

    s = _slide(prs, "GHG emissions by scope and quarter (tCO2e)")
    cd = CategoryChartData()
    cd.categories = QUARTERS
    for name, key in (("Scope 1", "scope1"), ("Scope 2", "scope2"), ("Scope 3", "scope3")):
        cd.add_series(name, [round(q[x][key], 1) for x in QUARTERS])
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.6), Inches(1.4), Inches(8), Inches(5.6), cd).chart
    ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
    peak = max(QUARTERS, key=lambda x: q[x]["scope1"] + q[x]["scope2"] + q[x]["scope3"])
    _bullets(s, [f"Scope 2 grid electricity is {100 * y['scope2'] / total:.0f}% of the footprint.",
                 f"{peak} is the peak quarter: summer cooling load at stores.",
                 f"Scope 3 covers upstream freight and business air travel, {y['scope3']:,.0f} tCO2e."], top=1.6, size=15)
    for shape in s.shapes:
        if shape.has_text_frame and shape.shape_id != s.shapes.title.shape_id:
            shape.left, shape.width = Inches(8.9), Inches(4.1)

    s = _slide(prs, "Energy, renewable share and waste diversion by quarter (%)")
    cd = CategoryChartData()
    cd.categories = QUARTERS
    cd.add_series("Renewable electricity share", [round(q[x]["renewable_share_pct"], 1) for x in QUARTERS])
    cd.add_series("Waste diversion", [round(q[x]["waste_diversion_pct"], 1) for x in QUARTERS])
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.4), Inches(8), Inches(5.6), cd).chart
    ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
    _bullets(s, [f"Energy use {y['energy_gj']:,.0f} GJ.",
                 f"Renewable share {y['renewable_share_pct']:.1f}% against a 12% target: rooftop solar exists only at "
                 "warehouses and offices.",
                 f"Waste diversion {y['waste_diversion_pct']:.1f}% against a 65% target."], top=1.6, size=15)
    for shape in s.shapes:
        if shape.has_text_frame and shape.shape_id != s.shapes.title.shape_id:
            shape.left, shape.width = Inches(8.9), Inches(4.1)

    s = _slide(prs, "Data quality: what was checked and fixed")
    counts = Counter(i.check for i in issues)
    adjusted = sum(r["status"] != "matched" for r in recon)
    _bullets(s, [f"{n_submissions} submissions received for {n_expected} expected site months.",
                 f"{counts['missing_submission']} missing and {counts['duplicate_submission']} duplicate submissions, "
                 f"{counts['unit_error']} unit errors (MWh entered as kWh), {counts['negative_value']} negative reading, "
                 f"{counts['recycled_exceeds_total']} rows with recycled waste above total waste.",
                 f"Electricity reconciled with utility bills for {len(recon)} site quarters: {adjusted} under reported by "
                 "5 to 7% and aligned to the bill.",
                 "Every correction and estimate is logged in the workbook's Validation Log for audit."])

    s = _slide(prs, "Risks and improvement opportunities")
    top = ", ".join(site for site, _ in inv["top_sites"][:3])
    _bullets(s, [f"Highest Scope 1 and 2 sites: {top}. Prioritise energy audits and LED and HVAC upgrades there.",
                 "Renewable share target missed: extend rooftop solar or green power purchase to large format stores.",
                 "Intensity reduction below target: store cooling drives Q1 peaks; set temperature and opening hour SOPs.",
                 "Recurring data risks: late submissions and unit slips. Add input range checks at source and a "
                 "monthly bill upload."])

    s = _slide(prs, "Methodology and sources")
    _bullets(s, ["GHG Protocol Corporate Standard, operational control; Scope 2 location based.",
                 "Grid factor 0.727 kg CO2/kWh, CEA CO2 Baseline Database Version 20.0 (FY 2023-24).",
                 "Diesel and LPG from IPCC 2006 defaults; R-410A GWP 1,924 (IPCC AR5).",
                 "Scope 3 freight and air travel from DEFRA 2024 conversion factors.",
                 "Disclosures mapped to SEBI BRSR Principle 6 and GRI 302, 303, 305 and 306."], size=15)
    prs.save(path)
    return path
